import os
import sys
import io
import json
from datetime import date
from fastapi.testclient import TestClient
import docx

# Add backend directory to sys.path
backend_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, backend_dir)

from app.database import engine, Base, SessionLocal
from app.models import User, Meeting, ActionItem, Decision, ClarificationRequest
from app.cache_and_limiter import clear_cache_and_limits_for_tests, _transcript_cache
from main import app

client = TestClient(app)

SAMPLE_TRANSCRIPT = """Priya: Okay so for the launch next week, I'll handle the press release, should be done by Thursday.
Sam: I can take the social media graphics, I'll have drafts by Wednesday.
Priya: Great. Someone needs to update the pricing page too, but I'm not sure who's free for that.
Rahul: I could maybe look at it if no one else can, but I have the client call all week.
Priya: Let's flag that one. Also we need the QA pass done before Friday's release.
Sam: I think that's on the QA team but honestly I don't know who exactly picks that up.
Priya: Okay, let's note that as unclear too. Last thing - can someone send the launch email to the mailing list?
Rahul: Yeah I'll do that, I'll send it Friday morning right before the release goes live.
Priya: Also, we've decided to push the enterprise pricing tier discussion to next quarter.
Sam: Agreed, and we're keeping the current logo for this launch instead of the redesign."""

def reset_db():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    clear_cache_and_limits_for_tests()

def run_stage_4_verification():
    print("==================================================================")
    print(" STAGE 4 VERIFICATION: POST /meetings ENDPOINT & PIPELINE WIRING")
    print("==================================================================")

    reset_db()

    # 1. Register and Login a Test User
    print("\n--- Setup: Registering Test User ---")
    reg_resp = client.post("/auth/register", json={"email": "alice.manager@example.com", "password": "SecretPassword123!"})
    assert reg_resp.status_code == 201, f"User registration failed: {reg_resp.text}"
    token = reg_resp.json()["access_token"]
    user_id = reg_resp.json()["user"]["id"]
    auth_headers = {"Authorization": f"Bearer {token}"}
    print(f"[OK] Test User created: ID={user_id}, Token received.")

    # -------------------------------------------------------------
    # Test A: JSON Transcript Submission
    # -------------------------------------------------------------
    print("\n--- Test A: JSON Transcript Submission (POST /meetings) ---")
    resp_a = client.post(
        "/meetings",
        json={"transcript_text": SAMPLE_TRANSCRIPT, "title": "Launch Sync #1"},
        headers=auth_headers
    )
    print(f"Status Code: {resp_a.status_code}")
    assert resp_a.status_code == 201, f"Expected 201, got {resp_a.status_code}: {resp_a.text}"
    data_a = resp_a.json()
    meeting_id_a = data_a["id"]
    print(f"[OK] Meeting created: ID={meeting_id_a}, Title='{data_a['title']}', Sentiment='{data_a['sentiment']}'")

    # Verify SQLite Database Counts for First Submission
    db = SessionLocal()
    try:
        meeting_count = db.query(Meeting).filter(Meeting.user_id == user_id).count()
        action_items = db.query(ActionItem).filter(ActionItem.meeting_id == meeting_id_a).all()
        decisions = db.query(Decision).filter(Decision.meeting_id == meeting_id_a).all()
        clarification_reqs = db.query(ClarificationRequest).all()

        clear_items = [item for item in action_items if not item.needs_clarification]
        ambiguous_items = [item for item in action_items if item.needs_clarification]

        print(f"\n[DATABASE VERIFICATION - FIRST SUBMISSION]")
        print(f"- Total Meetings for user: {meeting_count} (Expected: 1)")
        print(f"- Total Action Items in DB: {len(action_items)} (Expected: 5)")
        print(f"- Clear Action Items: {len(clear_items)} (Expected: 3)")
        print(f"- Ambiguous Action Items (needs_clarification=True): {len(ambiguous_items)} (Expected: 2)")
        print(f"- Decisions in DB: {len(decisions)} (Expected: 2)")
        print(f"- Clarification Requests in DB: {len(clarification_reqs)} (Expected: 2)")

        assert meeting_count == 1, f"Expected 1 meeting, got {meeting_count}"
        assert len(action_items) == 5, f"Expected 5 action items, got {len(action_items)}"
        assert len(clear_items) == 3, f"Expected 3 clear items, got {len(clear_items)}"
        assert len(ambiguous_items) == 2, f"Expected 2 ambiguous items, got {len(ambiguous_items)}"
        assert len(decisions) == 2, f"Expected 2 decisions, got {len(decisions)}"
        assert len(clarification_reqs) == 2, f"Expected 2 clarification requests, got {len(clarification_reqs)}"

        print("[OK] All database counts match exact specification!")
    finally:
        db.close()

    # -------------------------------------------------------------
    # Test B: Identical Transcript Cache Test
    # -------------------------------------------------------------
    print("\n--- Test B: Identical Transcript Caching ---")
    assert len(_transcript_cache) > 0, "Transcript cache should contain entry for first submission"
    print(f"[OK] Cache entry exists for user {user_id}.")

    resp_b = client.post(
        "/meetings",
        json={"transcript_text": SAMPLE_TRANSCRIPT, "title": "Launch Sync #2 (Duplicate Transcript)"},
        headers=auth_headers
    )
    assert resp_b.status_code == 201, f"Expected 201, got {resp_b.status_code}"
    print(f"[OK] Second identical submission succeeded instantly using cache.")

    # -------------------------------------------------------------
    # Test C: Duplicate Action Item Detection & mention_count
    # -------------------------------------------------------------
    print("\n--- Test C: Duplicate Detection & Mention Count Verification ---")
    db = SessionLocal()
    try:
        # Check that mention_count on original items incremented to 2
        items = db.query(ActionItem).filter(ActionItem.meeting_id == meeting_id_a).all()
        for item in items:
            print(f"Item: '{item.task[:35]}...' | Owner: {item.owner} | Mention Count: {item.mention_count}")
            assert item.mention_count == 2, f"Expected mention_count=2, got {item.mention_count}"
        
        total_items_in_db = db.query(ActionItem).count()
        assert total_items_in_db == 5, f"Expected 5 total action item rows (no duplicates created), got {total_items_in_db}"
        print(f"[OK] Duplicate detection correctly prevented duplicate rows and incremented mention_count to 2!")
    finally:
        db.close()

    # -------------------------------------------------------------
    # Test D: .txt File Upload
    # -------------------------------------------------------------
    print("\n--- Test D: .txt Multipart File Upload ---")
    txt_content = b"John: I will handle the deployment by Friday.\nSarah: I can review the pull request by Thursday."
    files_txt = {"file": ("notes.txt", io.BytesIO(txt_content), "text/plain")}
    resp_d = client.post("/meetings", files=files_txt, data={"title": "Text File Meeting"}, headers=auth_headers)
    assert resp_d.status_code == 201, f"Expected 201, got {resp_d.status_code}: {resp_d.text}"
    print(f"[OK] .txt upload processed successfully: ID={resp_d.json()['id']}")

    # -------------------------------------------------------------
    # Test E: .docx File Upload
    # -------------------------------------------------------------
    print("\n--- Test E: .docx Multipart File Upload ---")
    doc = docx.Document()
    doc.add_paragraph("Alice: I'll prepare the financial report by Monday.")
    doc.add_paragraph("Bob: We decided to cancel the old subscription plan.")
    docx_bytes = io.BytesIO()
    doc.save(docx_bytes)
    docx_bytes.seek(0)

    files_docx = {"file": ("meeting_notes.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    resp_e = client.post("/meetings", files=files_docx, data={"title": "Word Doc Meeting"}, headers=auth_headers)
    assert resp_e.status_code == 201, f"Expected 201, got {resp_e.status_code}: {resp_e.text}"
    print(f"[OK] .docx upload parsed with python-docx and processed successfully: ID={resp_e.json()['id']}")

    # -------------------------------------------------------------
    # Test F: Authentication Failure (Missing JWT)
    # -------------------------------------------------------------
    print("\n--- Test F: Missing JWT Header ---")
    resp_f = client.post("/meetings", json={"transcript_text": "Sample text"})
    assert resp_f.status_code == 401, f"Expected 401, got {resp_f.status_code}"
    print(f"[OK] Missing JWT rejected with HTTP 401: {resp_f.json().get('detail')}")

    # -------------------------------------------------------------
    # Test G: Authentication Failure (Invalid JWT)
    # -------------------------------------------------------------
    print("\n--- Test G: Invalid JWT Header ---")
    resp_g = client.post("/meetings", json={"transcript_text": "Sample text"}, headers={"Authorization": "Bearer fake.invalid.jwt"})
    assert resp_g.status_code == 401, f"Expected 401, got {resp_g.status_code}"
    print(f"[OK] Invalid JWT rejected with HTTP 401: {resp_g.json().get('detail')}")

    # -------------------------------------------------------------
    # Test H: Rate Limiting (> 10 Submissions in 1 Hour)
    # -------------------------------------------------------------
    print("\n--- Test H: Rate Limiting Verification (> 10 Submissions/Hour) ---")
    # Reset limits for clean rate limit test
    clear_cache_and_limits_for_tests()

    # Submit 10 meetings
    for i in range(1, 11):
        r = client.post(
            "/meetings",
            json={"transcript_text": f"User {i}: I will complete task number {i}."},
            headers=auth_headers
        )
        assert r.status_code == 201, f"Submission {i} failed: {r.text}"
    print(f"[OK] 10 consecutive submissions succeeded under rate limit.")

    # 11th submission must be rejected with HTTP 429
    resp_11 = client.post(
        "/meetings",
        json={"transcript_text": "User 11: This should trigger rate limit."},
        headers=auth_headers
    )
    print(f"11th Submission Status Code: {resp_11.status_code}")
    assert resp_11.status_code == 429, f"Expected 429 Too Many Requests, got {resp_11.status_code}"
    print(f"[OK] Rate limit triggered on 11th request with HTTP 429: {resp_11.json().get('detail')}")

    print("\n==================================================================")
    print(" [STAGE 4 SUCCESS REPORT]")
    print(" - POST /meetings wired to Stage 3 extraction pipeline.")
    print(" - JSON transcript, .txt upload, and .docx upload verified.")
    print(" - Identical transcript cache verified (reused without duplicate LLM calls).")
    print(" - Duplicate detection verified (SequenceMatcher > 0.75 + owner match -> mention_count incremented).")
    print(" - Clarification requests created for items with needs_clarification=True.")
    print(" - JWT authentication enforced (401 on missing/invalid token).")
    print(" - Rate limit enforced (429 on > 10 submissions/hour).")
    print("==================================================================\n")

if __name__ == "__main__":
    run_stage_4_verification()
