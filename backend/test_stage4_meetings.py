import os
import sys
import io
import json
from datetime import date
import docx

# Add backend directory to sys.path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi.testclient import TestClient
from app.database import Base, engine, SessionLocal
from app.models import User, Meeting, ActionItem, Decision, ClarificationRequest
from main import app

client = TestClient(app)

SAMPLE_TRANSCRIPT = """Priya: Okay so for the launch next week, I'll handle the press release, should be done by Thursday.
Sam: I can take the social media graphics, I'll have drafts by Wednesday.
Priya: Great. Someone needs to update the pricing page too, but I'm not sure who's free for that.
Rahul: I could maybe look at it if no one else can, but I have the client call all week.
Priya: Let's flag that one. Also we need the QA pass done before Friday's release.
Sam: I think that's on the QA team but honestly I don't know who exactly picks that up.
Priya: Okay, let's note that as unclear too. Last thing - can someone send the launch email to the mailing list?
Rahul: Yeah I'll do that, I'll send it Friday morning right before the release goes live.
Priya: Also, we've decided to push the enterprise pricing tier discussion to next quarter.
Sam: Agreed, and we're keeping the current logo for this launch instead of the redesign."""

def run_stage4_tests():
    print("==================================================================")
    print("=== Stage 4: Testing POST /meetings & Pipeline Integration ===")
    print("==================================================================")

    # 1. Initialize fresh DB
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    print("[INIT] Fresh database initialized.")

    # 2. Register Users (User 1 and User 2)
    res_u1 = client.post("/auth/register", json={
        "email": "sarah.lead@clarifier.com",
        "password": "Password123!",
        "role": "member"
    })
    assert res_u1.status_code == 201
    token1 = res_u1.json()["access_token"]
    user1_id = res_u1.json()["user"]["id"]
    auth_headers1 = {"Authorization": f"Bearer {token1}"}

    res_u2 = client.post("/auth/register", json={
        "email": "other.user@clarifier.com",
        "password": "Password123!",
        "role": "member"
    })
    assert res_u2.status_code == 201
    token2 = res_u2.json()["access_token"]
    user2_id = res_u2.json()["user"]["id"]
    auth_headers2 = {"Authorization": f"Bearer {token2}"}
    print(f"[OK] Test users registered: User 1 (ID {user1_id}), User 2 (ID {user2_id})")

    # -------------------------------------------------------------
    # TEST 1: Unauthenticated request to POST /meetings -> 401
    # -------------------------------------------------------------
    print("\n--- [TEST 1] Unauthenticated Request Rejection (401) ---")
    unauth_res = client.post("/meetings", json={"transcript_text": SAMPLE_TRANSCRIPT})
    print(f"Status: {unauth_res.status_code}")
    assert unauth_res.status_code == 401, f"Expected 401, got {unauth_res.status_code}"
    print("[PASS] Test 1: Missing JWT correctly returned 401 Unauthorized.")

    # -------------------------------------------------------------
    # TEST 2: Empty Transcript Rejection -> 400
    # -------------------------------------------------------------
    print("\n--- [TEST 2] Empty Transcript Rejection (400) ---")
    empty_res = client.post("/meetings", json={"transcript_text": "   "}, headers=auth_headers1)
    print(f"Status: {empty_res.status_code}")
    assert empty_res.status_code == 400, f"Expected 400, got {empty_res.status_code}"
    print("[PASS] Test 2: Empty transcript correctly returned 400 Bad Request.")

    # -------------------------------------------------------------
    # TEST 3: Unsupported File Upload -> 400
    # -------------------------------------------------------------
    print("\n--- [TEST 3] Unsupported File Format Rejection (400) ---")
    bad_file = io.BytesIO(b"dummy executable content")
    unsupported_res = client.post(
        "/meetings",
        files={"file": ("malicious.exe", bad_file, "application/octet-stream")},
        headers=auth_headers1
    )
    print(f"Status: {unsupported_res.status_code}, Detail: {unsupported_res.json()}")
    assert unsupported_res.status_code == 400, f"Expected 400, got {unsupported_res.status_code}"
    print("[PASS] Test 3: Unsupported file format correctly returned 400 Bad Request.")

    # -------------------------------------------------------------
    # TEST 4: JSON Submission of EXACT Sample Transcript
    # -------------------------------------------------------------
    print("\n--- [TEST 4] POST /meetings JSON with Sample Transcript ---")
    create_res = client.post(
        "/meetings",
        json={
            "title": "Launch Readiness Sync",
            "transcript_text": SAMPLE_TRANSCRIPT,
            "meeting_date": date.today().isoformat()
        },
        headers=auth_headers1
    )
    print(f"Status: {create_res.status_code}")
    assert create_res.status_code == 201, f"Expected 201, got {create_res.status_code}"
    m_data = create_res.json()
    meeting1_id = m_data["id"]
    print(f"Created Meeting ID: {meeting1_id}, Title: {m_data['title']}, Sentiment: {m_data['sentiment']}")
    assert m_data["user_id"] == user1_id
    assert m_data["transcript_text"] == SAMPLE_TRANSCRIPT
    assert len(m_data["action_items"]) == 5, f"Expected 5 action items, got {len(m_data['action_items'])}"
    assert len(m_data["decisions"]) == 2, f"Expected 2 decisions, got {len(m_data['decisions'])}"

    # Verify SQLite DB State for Meeting 1
    db = SessionLocal()
    try:
        db_items = db.query(ActionItem).filter(ActionItem.meeting_id == meeting1_id).all()
        assert len(db_items) == 5
        
        clarification_items = [i for i in db_items if i.needs_clarification]
        assert len(clarification_items) == 2, f"Expected 2 clarification items, got {len(clarification_items)}"
        
        db_clar_reqs = db.query(ClarificationRequest).all()
        assert len(db_clar_reqs) == 2, f"Expected exactly 2 clarification_requests, got {len(db_clar_reqs)}"
        for cr in db_clar_reqs:
            assert cr.answered_at is None
            assert cr.reminder_sent_at is None
            assert cr.question_sent_at is not None

        db_decisions = db.query(Decision).filter(Decision.meeting_id == meeting1_id).all()
        assert len(db_decisions) == 2
        print(f"[OK] Verified Meeting 1 in SQLite: 5 action items, 2 clarification requests, 2 decisions.")
    finally:
        db.close()
    print("[PASS] Test 4: JSON submission processed and persisted accurately.")

    # -------------------------------------------------------------
    # TEST 5: Duplicate Submission (Same user submits same transcript)
    # -------------------------------------------------------------
    print("\n--- [TEST 5] Duplicate Submission Handling (Mention Count Increment) ---")
    dup_res = client.post(
        "/meetings",
        json={
            "title": "Launch Readiness Sync - Followup",
            "transcript_text": SAMPLE_TRANSCRIPT,
            "meeting_date": date.today().isoformat()
        },
        headers=auth_headers1
    )
    assert dup_res.status_code == 201
    m2_data = dup_res.json()
    meeting2_id = m2_data["id"]
    assert meeting2_id != meeting1_id

    # Verify DB State: Meeting 2 exists, but action items were NOT duplicated; mention_count incremented
    db = SessionLocal()
    try:
        total_meetings_u1 = db.query(Meeting).filter(Meeting.user_id == user1_id).count()
        assert total_meetings_u1 == 2, f"Expected 2 meetings for user 1, got {total_meetings_u1}"

        total_items_u1 = (
            db.query(ActionItem)
            .join(Meeting, ActionItem.meeting_id == Meeting.id)
            .filter(Meeting.user_id == user1_id)
            .all()
        )
        assert len(total_items_u1) == 5, f"Expected 5 action items (no duplicates), found {len(total_items_u1)}"
        
        # Verify mention_count = 2 for all 5 items
        for item in total_items_u1:
            assert item.mention_count == 2, f"Expected mention_count=2, got {item.mention_count} for task '{item.task}'"
        print(f"[OK] Verified duplicate submission: 2 meetings stored, 0 duplicate action item rows created, all 5 items have mention_count=2.")
    finally:
        db.close()
    print("[PASS] Test 5: Duplicate detection preserved data integrity and incremented mention_count.")

    # -------------------------------------------------------------
    # TEST 6: User Isolation Check (User 2 submits same transcript)
    # -------------------------------------------------------------
    print("\n--- [TEST 6] Cross-User Isolation (User 2 Submits Transcript) ---")
    u2_res = client.post(
        "/meetings",
        json={
            "title": "User 2 Launch Meeting",
            "transcript_text": SAMPLE_TRANSCRIPT,
            "meeting_date": date.today().isoformat()
        },
        headers=auth_headers2
    )
    assert u2_res.status_code == 201
    
    db = SessionLocal()
    try:
        u2_items = (
            db.query(ActionItem)
            .join(Meeting, ActionItem.meeting_id == Meeting.id)
            .filter(Meeting.user_id == user2_id)
            .all()
        )
        assert len(u2_items) == 5, f"Expected 5 new items for User 2, got {len(u2_items)}"
        for item in u2_items:
            assert item.mention_count == 1, "User 2's action items should start at mention_count=1"
        print("[OK] User 2 items created independently with mention_count=1; never merged across users.")
    finally:
        db.close()
    print("[PASS] Test 6: Cross-user data isolation verified.")

    # -------------------------------------------------------------
    # TEST 7: Plain Text (.txt) File Upload
    # -------------------------------------------------------------
    print("\n--- [TEST 7] Multipart .txt File Upload ---")
    txt_content = b"Alice: I'll prepare the financial report by Monday.\nBob: We decided to cancel the old subscription."
    txt_file = io.BytesIO(txt_content)
    txt_res = client.post(
        "/meetings",
        files={"file": ("meeting_notes.txt", txt_file, "text/plain")},
        data={"title": "Finance Review"},
        headers=auth_headers1
    )
    print(f"Status: {txt_res.status_code}")
    assert txt_res.status_code == 201, f"Expected 201, got {txt_res.status_code}"
    txt_data = txt_res.json()
    assert txt_data["title"] == "Finance Review"
    assert "financial report" in txt_data["transcript_text"]
    print("[PASS] Test 7: Plain text file upload processed successfully.")

    # -------------------------------------------------------------
    # TEST 8: Word Document (.docx) File Upload
    # -------------------------------------------------------------
    print("\n--- [TEST 8] Multipart .docx File Upload ---")
    doc = docx.Document()
    doc.add_heading("Engineering Sync", 0)
    doc.add_paragraph("John: I will complete task number 1 by Friday.")
    doc.add_paragraph("Sarah: I will complete task number 2 by Friday.")
    docx_stream = io.BytesIO()
    doc.save(docx_stream)
    docx_stream.seek(0)

    docx_res = client.post(
        "/meetings",
        files={"file": ("eng_sync.docx", docx_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=auth_headers1
    )
    print(f"Status: {docx_res.status_code}")
    assert docx_res.status_code == 201, f"Expected 201, got {docx_res.status_code}"
    docx_data = docx_res.json()
    assert "task number 1" in docx_data["transcript_text"]
    print("[PASS] Test 8: Word document (.docx) upload and text extraction processed successfully.")

    print("\n==================================================================")
    print("=== ALL STAGE 4 TESTS PASSED SUCCESSFULLY ===")
    print("==================================================================")
    return True

if __name__ == "__main__":
    run_stage4_tests()
